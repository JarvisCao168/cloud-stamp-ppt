#!/usr/bin/env python3
import os, json, random
from collections import Counter
from pptx import Presentation
from pptx.util import Inches

TEMPLATE_DIR = r"F:\ppt模板"
OUTPUT_FILE = r"D:\yzppt\docs\template_analysis.json"
SAMPLES_PER_CATEGORY = 5

TARGET_CATEGORIES = [
    "国赛", "工作汇报", "莫兰迪PPT", "毕业答辩", "高端商务",
    "中国风格", "项目策划", "清新文艺", "扁平风格", "极简风格",
    "企业宣传", "简历求职", "教师课件", "培训课件", "述职报告",
]

def get_main_colors(slide):
    colors = []
    for shape in slide.shapes:
        if hasattr(shape, "fill"):
            try:
                if shape.fill.type == 1:  # SOLID
                    color = shape.fill.fore_color.rgb
                    if color:
                        colors.append(str(color))
            except: pass
        if hasattr(shape, "line"):
            try:
                if shape.line.fill.type == 1:
                    color = shape.line.fill.fore_color.rgb
                    if color:
                        colors.append(str(color))
            except: pass
    return colors[:10]

def get_fonts(slide):
    fonts = []
    for shape in slide.shapes:
        if hasattr(shape, "text_frame"):
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    f = run.font
                    if f.name:
                        fonts.append({"name": f.name, "size": f.size.pt if f.size else None})
    return fonts[:10]

def analyze_template(filepath):
    try:
        prs = Presentation(filepath)
        slides_info = []
        all_colors = []
        all_fonts = []
        
        for slide in prs.slides:
            slide_info = {
                "layout": slide.slide_layout.name if slide.slide_layout else "Unknown",
                "colors": get_main_colors(slide),
                "fonts": get_fonts(slide),
            }
            slides_info.append(slide_info)
            all_colors.extend(slide_info["colors"])
            all_fonts.extend(slide_info["fonts"])
        
        color_counts = Counter(all_colors)
        dominant_colors = [c for c, _ in color_counts.most_common(5)]
        
        font_counts = Counter([f["name"] for f in all_fonts if f.get("name")])
        dominant_fonts = [{"name": n, "count": c} for n, c in font_counts.most_common(5)]
        
        return {
            "filename": os.path.basename(filepath),
            "path": filepath,
            "total_slides": len(prs.slides),
            "dominant_colors": dominant_colors,
            "dominant_fonts": dominant_fonts,
            "slides": slides_info[:8],
        }
    except Exception as e:
        return None

def main():
    samples = []
    for cat in TARGET_CATEGORIES:
        cat_path = os.path.join(TEMPLATE_DIR, cat)
        if not os.path.exists(cat_path): continue
        
        ppt_files = []
        for root, dirs, files in os.walk(cat_path):
            for f in files:
                if f.endswith(".pptx"):
                    ppt_files.append(os.path.join(root, f))
        
        if not ppt_files: continue
        selected = random.sample(ppt_files, min(SAMPLES_PER_CATEGORY, len(ppt_files)))
        for fp in selected:
            samples.append({"category": cat, "filepath": fp})
    
    print(f"选择 {len(samples)} 个模板进行分析...")
    results = []
    for i, s in enumerate(samples, 1):
        print(f"[{i}/{len(samples)}] {s['category']}/{s['filepath'].split(chr(92))[-1]}")
        r = analyze_template(s["filepath"])
        if r:
            r["category"] = s["category"]
            results.append(r)
    
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    
    print(f"\n成功分析 {len(results)} 个模板")
    print(f"保存到 {OUTPUT_FILE}")
    
    # 汇总统计
    all_colors = []
    all_fonts = []
    for r in results:
        all_colors.extend(r["dominant_colors"])
        all_fonts.extend([f["name"] for f in r["dominant_fonts"] if f.get("name")])
    
    print("\n最常见颜色:")
    for c, n in Counter(all_colors).most_common(10):
        print(f"  #{c}: {n}次")
    print("\n最常见字体:")
    for f, n in Counter(all_fonts).most_common(10):
        print(f"  {f}: {n}次")

if __name__ == "__main__":
    main()
